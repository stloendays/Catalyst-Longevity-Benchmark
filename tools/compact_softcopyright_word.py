from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


SHORT_DESCRIPTIONS = [
    "长期稳定性概览与决策状态",
    "实验数据、原始记录与质量检查",
    "T95/T90/T80 寿命区间",
    "共同观测时间性能对比",
    "实验建议、可执行性与依据",
    "公开事实与论文实验窗口",
    "资料关联、确认与追溯",
    "AI 可用资料与使用边界",
    ".clrproj 项目管理",
    "版本、平台、格式与功能信息",
    "条件不一致时阻止直接排名",
]


def set_cell_margins(cell, value: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def compact_document(path: Path) -> None:
    document = Document(path)
    if len(document.tables) < 2:
        raise RuntimeError("Expected cover information table and interface summary table")

    summary = document.tables[1]
    if len(summary.rows) != 12:
        raise RuntimeError(f"Expected 12 rows in interface summary table, got {len(summary.rows)}")

    for row_index, row in enumerate(summary.rows):
        for column_index, cell in enumerate(row.cells):
            set_cell_margins(cell, 28 if row_index else 38)
            if row_index > 0 and column_index == 2:
                cell.text = SHORT_DESCRIPTIONS[row_index - 1]
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 0.92
                if column_index < 2:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(8.2 if row_index > 0 else 8.8)
                    run.font.name = "Microsoft YaHei"
                    rpr = run._element.get_or_add_rPr()
                    if rpr.rFonts is not None:
                        for key in ("eastAsia", "ascii", "hAnsi"):
                            rpr.rFonts.set(qn(f"w:{key}"), "Microsoft YaHei")

    prefixes = (
        "智策是一套面向催化剂长期稳定性研究",
        "本组截图由 Windows 原生程序",
        "截图中的内置数据",
    )
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefixes):
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.08
            for run in paragraph.runs:
                run.font.size = Pt(10)

    document.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    args = parser.parse_args()
    compact_document(args.docx)


if __name__ == "__main__":
    main()
