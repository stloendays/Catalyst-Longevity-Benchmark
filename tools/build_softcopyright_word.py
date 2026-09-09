from __future__ import annotations

import argparse
import tempfile
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SCREENSHOTS = [
    (
        "首页 - 长期稳定性总览",
        "01_首页_长期稳定性总览.png",
        "展示长期稳定性数据概览、催化剂数量、实验状态、条件检查与当前实验决策信息。",
    ),
    (
        "数据 - 质量检查",
        "02_数据_质量检查.png",
        "展示内置或导入实验数据、原始记录表以及数据质量检查状态。",
    ),
    (
        "分析 - T90 寿命指标",
        "03_分析_T90寿命指标.png",
        "展示 T95、T90、T80 等寿命阈值及区间/删失状态，避免把未实测区间插值为伪精确寿命。",
    ),
    (
        "分析 - 同时间对比",
        "04_分析_同时间对比.png",
        "展示不同催化剂在共同实际观测时间下的性能对比，并结合实验条件可比性决定是否输出领先判断。",
    ),
    (
        "分析 - 实验建议",
        "05_分析_实验建议.png",
        "展示后续实验建议、排期约束、可执行性、原因、下一步动作和建议依据。",
    ),
    (
        "分析 - 公开参考库",
        "06_分析_公开参考库.png",
        "展示内置公开事实与论文实验窗口，并用于核对实验条件和建议量级。",
    ),
    (
        "资料 - 关联与确认",
        "07_资料_关联与确认.png",
        "展示资料来源、识别内容、催化剂与时间关联、人工确认状态及可追溯信息。",
    ),
    (
        "AI 助手 - 资料准备",
        "08_AI助手_资料准备.png",
        "展示 AI 可用资料边界、已确认资料状态以及受控分析流程；未确认资料不会自动进入 AI 上下文。",
    ),
    (
        "项目 - 本地管理",
        "09_项目_本地管理.png",
        "展示 .clrproj 本地项目的新建、打开、保存、另存为以及项目状态管理。",
    ),
    (
        "设置 - 软件信息",
        "10_设置_软件信息.png",
        "展示软件名称、版本、运行平台、技术架构、输入格式、项目存储和核心功能等固定信息。",
    ),
    (
        "分析 - 条件差异阻止比较",
        "11_分析_条件差异阻止比较.png",
        "展示实验条件明确不一致时停止直接排名的保护逻辑，避免跨条件误比较。",
    ),
]

FONT = "Microsoft YaHei"


def set_run_font(run, size: float = 11, bold: bool = False) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), FONT)


def set_cell_shading(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    tcmar = tcpr.first_child_found_in("w:tcMar")
    if tcmar is None:
        tcmar = OxmlElement("w:tcMar")
        tcpr.append(tcmar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcmar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tcmar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_cell(cell, size=10.5, bold=False, center=False) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    for paragraph in cell.paragraphs:
        if center:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.15
        for run in paragraph.runs:
            set_run_font(run, size=size, bold=bold)


def add_heading(document: Document, text: str, size: float = 17, before: float = 0, after: float = 8) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=True)


def add_body(document: Document, text: str, size: float = 11, after: float = 8) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.3
    run = paragraph.add_run(text)
    set_run_font(run, size=size)


def compress_image(source: Path, destination: Path) -> None:
    image = Image.open(source).convert("RGB")
    image.save(destination, "JPEG", quality=92, optimize=True, progressive=True)


def build_document(screenshot_dir: Path, output: Path) -> None:
    for _, filename, _ in SCREENSHOTS:
        source = screenshot_dir / filename
        if not source.exists():
            raise FileNotFoundError(f"Missing screenshot: {source}")

    output.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)

    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    if "Figure Caption CN" not in document.styles:
        caption_style = document.styles.add_style("Figure Caption CN", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption_style = document.styles["Figure Caption CN"]
    caption_style.font.name = FONT
    caption_style.font.size = Pt(10.5)
    caption_style.font.color.rgb = RGBColor(0, 0, 0)
    caption_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    caption_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_style.paragraph_format.space_before = Pt(6)
    caption_style.paragraph_format.space_after = Pt(6)

    # Cover
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(105)
    set_run_font(paragraph.add_run("智策 V1.0"), size=26, bold=True)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(18)
    set_run_font(paragraph.add_run("催化剂长期稳定性评估与实验决策系统"), size=18, bold=True)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(40)
    set_run_font(paragraph.add_run("软件著作权界面材料（实际运行截图版）"), size=16)

    info = document.add_table(rows=5, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info.autofit = False
    info_rows = [
        ("软件简称", "智策"),
        ("版本号", "V1.0"),
        ("运行平台", "Windows 10/11 x64"),
        ("开发技术", "C++20、Qt 6、SQLite"),
        ("项目格式", ".clrproj（SQLite）"),
    ]
    for i, (key, value) in enumerate(info_rows):
        left, right = info.rows[i].cells
        left.text = key
        right.text = value
        set_cell_shading(left, "F2F2F2")
        style_cell(left, bold=True)
        style_cell(right)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(28)
    set_run_font(
        paragraph.add_run("本文件截图由最新 Windows Qt 程序自动生成，用于对应展示软件主要功能模块与操作界面。"),
        size=10,
    )

    # Overview
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    add_heading(document, "一、界面材料说明")
    add_body(document, "智策是一套面向催化剂长期稳定性研究的数据分析与实验决策桌面软件。软件主体采用 C++20 与 Qt 6 构建，并使用 SQLite 保存本地项目。")
    add_body(document, "本组截图由 Windows 原生程序自动载入内置演示场景后直接截取，覆盖数据检查、寿命指标、条件可比性、实验建议、资料确认、AI 使用边界、项目管理和软件信息等主要模块。")
    add_body(document, "截图中的内置数据用于软件功能演示和流程验证，不作为真实科研实验结论。")
    add_heading(document, "二、主要界面", before=12)

    summary = document.add_table(rows=1, cols=3)
    summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    summary.autofit = False
    for index, text in enumerate(("序号", "页面", "主要展示内容")):
        cell = summary.rows[0].cells[index]
        cell.text = text
        set_cell_shading(cell, "EDEDED")
        style_cell(cell, bold=True, center=index < 2)
    for index, (title, _, description) in enumerate(SCREENSHOTS, 1):
        cells = summary.add_row().cells
        cells[0].text = str(index)
        cells[1].text = title
        cells[2].text = description
        style_cell(cells[0], size=9.5, center=True)
        style_cell(cells[1], size=9.5, center=True)
        style_cell(cells[2], size=9.5)

    with tempfile.TemporaryDirectory() as temp_dir_name:
        temp_dir = Path(temp_dir_name)
        for index, (title, filename, description) in enumerate(SCREENSHOTS, 1):
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            add_heading(document, f"{index}. {title}", size=16, after=5)
            add_body(document, description, size=10.5, after=9)

            compressed = temp_dir / f"{index:02d}.jpg"
            compress_image(screenshot_dir / filename, compressed)
            picture_paragraph = document.add_paragraph()
            picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture_paragraph.paragraph_format.space_after = Pt(0)
            picture_paragraph.add_run().add_picture(str(compressed), width=Inches(7.0))

            caption = document.add_paragraph(style="Figure Caption CN")
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(caption.add_run(f"图 {index}  智策 - {title}（Windows 实际运行截图）"), size=10.5)

    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    add_heading(document, "三、申报一致性核对")
    checks = [
        "软件全称、简称和版本号应与申请表、操作说明书、安装程序和源代码材料保持一致。",
        "正式材料中的功能说明应以当前 V1.0 已实现的软件能力为准，不把尚未启用的外部服务描述为软件主体能力。",
        "截图不得包含 API Key、个人隐私、调试控制台或与登记无关的系统信息。",
        "用户实验数据、公开参考信息、系统计算结果和内置模拟数据应继续保持清晰的数据边界。",
    ]
    for i, text in enumerate(checks, 1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(7)
        set_run_font(paragraph.add_run(f"{i}. {text}"), size=11)

    check_table = document.add_table(rows=1, cols=3)
    check_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, text in enumerate(("核对项", "本材料口径", "正式申报要求")):
        cell = check_table.rows[0].cells[index]
        cell.text = text
        set_cell_shading(cell, "EDEDED")
        style_cell(cell, bold=True, center=True)
    rows = [
        ("软件全称", "催化剂长期稳定性评估与实验决策系统", "与申请表、说明书和源代码一致"),
        ("软件简称", "智策", "与程序标题栏及安装名称一致"),
        ("版本号", "V1.0", "与申请版本、EXE 和说明书一致"),
        ("主要技术", "C++20 / Qt 6 / SQLite", "与源代码和构建配置一致"),
        ("截图来源", "Windows Qt 实际运行截图", "保留当前自动截图生成记录"),
    ]
    for left_text, middle_text, right_text in rows:
        cells = check_table.add_row().cells
        cells[0].text = left_text
        cells[1].text = middle_text
        cells[2].text = right_text
        style_cell(cells[0], bold=True)
        style_cell(cells[1])
        style_cell(cells[2])

    props = document.core_properties
    props.title = "智策 V1.0 软件著作权界面材料"
    props.subject = "催化剂长期稳定性评估与实验决策系统"
    props.keywords = "智策, 软件著作权, 催化剂, Qt, C++"
    props.author = ""
    props.last_modified_by = ""
    document.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--screenshots", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    build_document(args.screenshots, args.output)
    print(args.output)


if __name__ == "__main__":
    main()
