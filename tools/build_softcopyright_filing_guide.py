from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

FULL_NAME = "催化剂长期稳定性评估与实验决策系统"
SHORT_NAME = "智策"
VERSION = "V1.0"
DEFAULT_OUTPUT = Path("docs/software-copyright/final/智策_V1.0_软著申报填写指南.docx")


def font(run, size=10.5, bold=None):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold


def shade(cell, fill="E7E6E6"):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def margins(cell, top=100, start=120, bottom=100, end=120):
    tcpr = cell._tc.get_or_add_tcPr()
    tcmar = tcpr.first_child_found_in("w:tcMar")
    if tcmar is None:
        tcmar = OxmlElement("w:tcMar")
        tcpr.append(tcmar)
    for name, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tcmar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tcmar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def width(cell, cm):
    tcpr = cell._tc.get_or_add_tcPr()
    tcw = tcpr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        tcpr.append(tcw)
    tcw.set(qn("w:w"), str(int(cm * 567)))
    tcw.set(qn("w:type"), "dxa")


def cell_style(cell, size=9.2, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    margins(cell)
    for p in cell.paragraphs:
        p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.1
        for r in p.runs:
            font(r, size=size, bold=bold)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    trpr.append(node)


def table(doc, headers, rows, widths, size=9.0):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    repeat_header(t.rows[0])
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        shade(c)
        width(c, widths[i])
        cell_style(c, size=size, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            width(cells[i], widths[i])
            align = WD_ALIGN_PARAGRAPH.CENTER if (i == 0 or (len(headers) >= 3 and i == 2)) else WD_ALIGN_PARAGRAPH.LEFT
            cell_style(cells[i], size=size, align=align)
        trpr = t.rows[-1]._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
    return t


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    font(r, size=15 if level == 1 else 12, bold=True)
    p.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    return p


def body(doc, text, size=10.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.35
    r = p.add_run(text)
    font(r, size=size)
    return p


def note(doc, title, text):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    c = t.cell(0, 0)
    width(c, 16.3)
    shade(c, "F2F2F2")
    margins(c, 140, 180, 140, 180)
    p = c.paragraphs[0]
    r = p.add_run(title)
    font(r, 10.2, True)
    p2 = c.add_paragraph()
    r2 = p2.add_run(text)
    font(r2, 9.8)
    p2.paragraph_format.line_spacing = 1.25
    p2.paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def copy_block(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label)
    font(r, 10.3, True)
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    c = t.cell(0, 0)
    c.text = text
    shade(c, "FAFAFA")
    width(c, 16.3)
    margins(c, 140, 180, 140, 180)
    cell_style(c, 9.8, align=WD_ALIGN_PARAGRAPH.JUSTIFY)


def page_number(p):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    inst = OxmlElement("w:instrText"); inst.set(qn("xml:space"), "preserve"); inst.text = " PAGE "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, inst, end])
    font(run, 8.5)


def build(output=DEFAULT_OUTPUT):
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.1); sec.right_margin = Cm(2.1)
    sec.header_distance = Cm(0.8); sec.footer_distance = Cm(0.8)
    props = doc.core_properties
    props.title = "智策 V1.0 软件著作权登记申报填写指南"
    props.subject = "软件著作权登记申报填写指南"
    props.author = ""; props.last_modified_by = ""
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"; normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(10.5)
    for name, sz in [("Heading 1", 15), ("Heading 2", 12)]:
        s = doc.styles[name]
        s.font.name = "Arial"; s._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        s.font.size = Pt(sz); s.font.bold = True; s.font.color.rgb = RGBColor(0,0,0)
    page_number(sec.footer.paragraphs[0])

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(85); p.paragraph_format.space_after = Pt(12)
    font(p.add_run("智策 V1.0"), 26, True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(24)
    font(p.add_run("软件著作权登记申报填写指南"), 19, True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(6)
    font(p.add_run(FULL_NAME), 12.5)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run("个人独立开发 · 原始取得 · 全部权利"), 11, True)
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(48); p.paragraph_format.space_after = Pt(4)
    font(p.add_run("使用说明"), 11, True)
    body(doc, "本指南按当前已冻结的“智策 V1.0”软件、源程序和说明材料编制。带“可直接填”的字段可按本文复制；带“本人确认”的字段涉及真实开发时间、发表情况或个人身份信息，提交前必须由申请人核实。", 10.2)
    note(doc, "重要", "本指南是申报操作辅助文件，不属于必须提交的程序/文档鉴别材料。中国版权保护中心页面字段名称或选项如有调整，以提交当日系统实际显示为准。")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    heading(doc, "1. 一页速填卡")
    table(doc, ["字段", "建议填写", "状态"], [
        ("软件全称", FULL_NAME, "可直接填"), ("软件简称", SHORT_NAME, "可直接填"), ("版本号", VERSION, "可直接填"),
        ("著作权人类型", "自然人", "可直接填"), ("开发方式", "独立开发", "可直接填"), ("权利取得方式", "原始取得", "可直接填"),
        ("权利范围", "全部权利", "可直接填"), ("软件类型", "应用软件（若系统提供该选项）", "按页面选项"),
        ("开发完成日期", "建议以 V1.0 实际冻结/完成日填写；当前候选：2026-09-09", "本人确认"),
        ("发表状态", "见第 5 节；当前 GitHub 仓库为 public，不能未经核实直接选“未发表”", "本人确认"),
        ("编程语言", "C++20", "可直接填"), ("源程序量", "约 6976 行（本次第一方源码非空行统计）", "可直接填"),
        ("运行平台", "Windows 10/11 x64", "可直接填"), ("数据库/本地存储", "SQLite；项目文件扩展名 .clrproj", "可直接填")
    ], [3.1, 10.7, 2.5], 9.0)
    note(doc, "一致性要求", f"申请表、源程序鉴别材料、软件操作说明书、界面截图中的软件名称和版本统一使用“{FULL_NAME} / {SHORT_NAME} / {VERSION}”。")

    heading(doc, "2. 基础登记信息逐项填写")
    table(doc, ["字段", "填写内容", "注意事项"], [
        ("软件全称", FULL_NAME, "不要再使用 Catalyst Longevity Research/Benchmark 作为对外登记名称。"),
        ("软件简称", SHORT_NAME, "若页面允许简称，统一填“智策”。"), ("版本号", VERSION, "所有上传材料保持 V1.0。"),
        ("软件作品说明", "原创软件", "如页面表述不同，选择与“独立原创/自主开发”最接近的选项。"),
        ("开发方式", "独立开发", "用户已确认该软件由本人个人独立完成。"), ("权利取得方式", "原始取得", "不是受让、继承或许可取得。"),
        ("权利范围", "全部权利", "以申请人对本人独立开发软件享有完整著作权为前提。"),
        ("著作权人", "【本人姓名】", "必须与身份证明完全一致，不使用网名或 GitHub 用户名。"),
        ("证件类型/号码", "【按本人证件填写】", "仅在官方申报页面填写；不要提交到公开 GitHub 仓库。"),
        ("联系地址/电话/邮箱", "【本人真实有效信息】", "用于受理、补正和证书联系。")
    ], [3.0, 5.5, 7.8], 9.0)

    heading(doc, "3. 技术信息填写")
    body(doc, "以下内容以当前仓库中已实现并通过 Windows CI 构建的 V1.0 为依据。硬件型号、内存容量等与个人电脑有关的字段，不应虚构；若页面要求具体数值，请按你的实际开发电脑补充。")
    table(doc, ["项目", "建议填写", "备注"], [
        ("开发语言", "C++20", "可直接填"), ("GUI 框架", "Qt 6 Widgets", "可直接填"),
        ("开发/构建工具", "CMake；Visual Studio 2022 / MSVC（当前 Windows 自动构建环境）", "如你的实际开发工具不同，以实际为准"),
        ("数据库", "SQLite", "可直接填"), ("Excel 读取", "QXlsx", "第三方依赖，不作为自有源码申报"),
        ("PDF 处理", "Qt PDF；QPdfWriter / QTextDocument", "可直接作为技术环境描述"), ("运行操作系统", "Windows 10/11 x64", "可直接填"),
        ("项目文件", ".clrproj（SQLite 本地项目文件）", "可直接填"), ("数据输入", "CSV、Excel .xlsx、内置演示数据集", "可直接填"),
        ("资料输入", "PDF、TXT、Markdown、CSV、TSV", "可直接填"), ("报告输出", "PDF", "可直接填"),
        ("第一方源程序量", "6976 个非空代码行；鉴别材料 60 页，每页 50 行", "与当前生成记录一致")
    ], [3.3, 8.1, 4.9], 8.9)

    heading(doc, "4. 可直接复制的文字说明")
    body(doc, "下面提供“短版”和“完整版”。如果申报页面有字数上限，优先使用短版；如允许较长描述，使用完整版。不要把“AI”写成软件唯一或主要功能，V1.0 的核心仍是本地科研数据分析与实验决策辅助。")
    copy_block(doc, "4.1 开发目的（短版）", "面向催化剂长期稳定性实验，提供实验数据整理、质量检查、寿命指标计算、条件一致性校验、同时间性能比较和后续实验建议，降低长期实验数据整理成本，提高比较结果的可复查性和实验安排效率。")
    copy_block(doc, "4.2 开发目的（完整版）", "本软件面向催化剂长期稳定性研究与实验管理场景，旨在将分散的实验记录、长期性能数据、实验条件和文献资料统一组织到本地桌面系统中。系统通过数据质量检查、寿命阈值计算、实验条件一致性校验、同时间比较和规则化实验建议，帮助研究人员减少人工整理工作，避免在实验条件不一致时进行不当横向排名，并将当前观测结果转化为可追溯的后续实验安排。")
    copy_block(doc, "4.3 主要功能（推荐版）", "软件支持 CSV 和 Excel 实验数据导入，自动完成重复时间点、异常数值、实验条件缺失和条件变化检查；绘制长期性能曲线并计算保持率、T95、T90、T80 等寿命指标；在跨催化剂比较前核对温度、空速、压力和进料条件，条件不一致时停止直接比较；自动寻找共同观测时间并进行性能对比；根据时间点数量、寿命状态和条件完整度生成下一步实验建议；支持论文和文本资料整理、催化剂与时间点关联、人工确认、受控 AI 资料范围、本地项目保存与恢复以及 PDF 报告导出。")
    copy_block(doc, "4.4 技术特点（推荐版）", "软件采用 C++20 与 Qt 6 Widgets 构建 Windows 原生桌面界面，使用 SQLite 保存本地 .clrproj 项目。分析逻辑强调观测数据边界：长期曲线仅连接实际观测点；寿命阈值无法精确确定时保留区间或下限表达；跨催化剂比较前执行实验条件一致性检查；资料需经过人工确认后才进入 AI 可用范围。系统支持 CSV/Excel 数据输入、PDF/TXT/Markdown 等资料输入，并可生成 PDF 报告，核心分析无需依赖外部在线服务。")
    copy_block(doc, "4.5 应用领域/行业（按页面最接近选项选择）", "科学研究和技术服务 / 化学与化工研发 / 催化剂研发与实验数据分析。若系统仅提供大类，优先选择与“科学研究和技术服务业”或“应用软件/科研辅助软件”最接近的类别。")

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    heading(doc, "5. 发表状态：提交前必须确认")
    note(doc, "当前已知事实", "GitHub 仓库目前为 public；GitHub API 显示仓库创建时间为 2026-09-05 17:51:24 UTC（北京时间/新加坡时间为 2026-09-06 01:51:24）。但仓库元数据本身不能证明它从创建时起就一直公开，也不能自动证明当时公开的内容已经构成“智策 V1.0”的首次发表。")
    body(doc, "建议按下面的事实判断，不要为了省事随意选择“未发表”：")
    table(doc, ["情形", "事实", "填写建议"], [
        ("A", "如果该仓库自创建起就是公开仓库，而且当时已经向公众提供了该软件的程序/可识别版本", "倾向按“已发表”填写；首次发表日期以首次真实向公众提供该软件的日期为准。"),
        ("B", "如果仓库创建时为私有，后来才改为公开", "以实际改为公开并向公众提供软件的日期作为候选首次发表日期。"),
        ("C", "如果公开仓库当时只有研究资料或早期不完整代码，V1.0 后来才首次公开", "应以 V1.0 首次实际公开的日期为候选，不机械使用仓库创建时间。"),
        ("D", "如果在正式申报前软件从未向公众提供，只在本人设备、私有环境或特定非公众范围使用", "可按真实情况选择“未发表”。")
    ], [1.2, 7.0, 8.1], 9.0)
    body(doc, "“发表权”是决定软件是否公之于众的权利。公开 GitHub 仓库可能涉及通过信息网络向公众提供软件，因此这里应按实际公开历史判断。若你不确定仓库何时从 private 变为 public，可以先核对 GitHub 仓库设置历史、通知邮件或最早可公开访问记录，再填写。", 9.8)

    heading(doc, "6. 鉴别材料与上传清单")
    table(doc, ["材料", "文件/内容", "状态", "说明"], [
        ("源程序鉴别材料", "智策_V1.0_源程序鉴别材料.docx", "已生成", "60 页；每页 50 行；前 30 页 + 后 30 页；第三方源码未纳入。"),
        ("软件文档鉴别材料", "智策_V1.0_软件操作说明书.docx", "已生成", "含真实 Windows Qt 界面与操作说明。"),
        ("界面材料", "智策_V1.0_软件著作权界面材料.docx", "已生成", "作为辅助留档/老师审核；是否单独上传以申报页面要求为准。"),
        ("申请信息核对表", "智策_V1.0_软件著作权申请信息核对表.docx", "已生成", "用于本人填写个人信息和最终核对，不作为法定鉴别材料替代申请表。"),
        ("身份证明", "本人有效身份证明", "本人准备", "仅提交官方要求渠道，不上传公开仓库。"),
        ("安装包/EXE", "智策 Windows 安装包或便携包", "可留存", "通常不是程序/文档鉴别材料的替代品；仅在受理端明确要求或内部审核时提供。")
    ], [3.2, 5.1, 2.2, 5.8], 8.7)
    note(doc, "现行鉴别材料规则", "《计算机软件著作权登记办法》要求提交软件著作权登记申请表、软件鉴别材料及相关证明文件；程序和文档鉴别材料通常由前、后各连续 30 页组成，程序每页通常不少于 50 行。当前“智策 V1.0”源程序材料已按 60 页、每页 50 行整理。")

    heading(doc, "7. 提交前 12 项最终核对")
    checks = [
        "□ 软件全称在申请表、源代码页眉、说明书封面和截图材料中完全一致。", "□ 软件简称统一为“智策”，版本号统一为“V1.0”。",
        "□ 著作权人姓名与身份证明完全一致。", "□ 开发方式选择“独立开发”，权利取得选择“原始取得”，权利范围选择“全部权利”。",
        "□ 开发完成日期由本人按真实 V1.0 完成日确认，不机械照抄建议日期。", "□ 发表状态与 GitHub 真实公开历史一致；如填“已发表”，首次发表日期也与证据一致。",
        "□ 源程序材料为本人第一方代码，未使用 Qt、QXlsx、SQLite 等第三方源码凑页数。", "□ 源程序前后页连续、页码连续、每页行数符合当前材料设定。",
        "□ 软件操作说明书中的功能均可在 V1.0 实际运行程序中找到。", "□ 内置模拟数据明确作为演示数据，不写成真实科研实验结论。",
        "□ 身份证号、手机号、家庭/联系地址等个人信息未提交到公开 GitHub 仓库。", "□ 提交前保存一份最终申请表 PDF/截图和全部上传文件的本地备份。"
    ]
    for text in checks:
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.3); p.paragraph_format.first_line_indent = Cm(-0.3); p.paragraph_format.space_after = Pt(3)
        font(p.add_run(text), 10.0)

    p = heading(doc, "8. 个人信息待填区"); p.paragraph_format.page_break_before = True
    table(doc, ["项目", "本人填写"], [
        ("著作权人姓名", "________________________________________"), ("证件类型", "________________________________________"),
        ("证件号码", "________________________________________"), ("联系电话", "________________________________________"), ("电子邮箱", "________________________________________"),
        ("联系地址", "________________________________________"), ("邮政编码", "________________________________________"), ("V1.0 实际开发完成日期", "________年____月____日"),
        ("发表状态", "□ 已发表    □ 未发表    （按真实公开历史确认）"), ("首次发表日期（如适用）", "________年____月____日"),
        ("首次发表地点/国家地区（如适用）", "________________________________________")
    ], [6.0, 10.3], 9.6)

    heading(doc, "9. 依据与文件位置")
    body(doc, "法规依据：国家版权局《计算机软件著作权登记办法》及《计算机软件保护条例》。官方软件登记办理机构为中国版权保护中心；实际在线填报字段和选项以提交当日登记系统为准。", 9.8)
    body(doc, "项目材料目录：docs/software-copyright/final/；软件界面截图目录：docs/software-copyright/screenshots/。", 9.8)
    body(doc, "本指南版本：智策 V1.0 申报辅助版。生成日期：2026-09-10。", 9.4)
    for p in doc.paragraphs:
        p.paragraph_format.widow_control = True
    doc.save(output)
    print(output)


if __name__ == "__main__":
    build()
